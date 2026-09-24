"""EngiTrace AI — Excel Import Service (Phase 4)

Provides controlled ingestion for EngiTrace_Template.xlsx:
  1. Parse & normalize workbook sheets
  2. Input validation & relationship integrity
  3. Deterministic compliance evaluation (dry-run)
  4. Safe transactional persistence with audit provenance logging
"""

import io
from typing import Dict, Any, List, Optional
from datetime import datetime
import openpyxl
from sqlalchemy.orm import Session
from sqlalchemy.exc import IntegrityError

from backend.app.domain.models import DigitalThread, AuditAction
from backend.app.models.db_models import (
    DBRequirement, DBRisk, DBControl, DBVerification, DBEvidence
)
from backend.app.schemas.excel_import import ImportExcelResponse
from backend.app.schemas.findings import FindingResponse
from backend.app.services.compliance_service import ComplianceService
from backend.app.services.audit_service import AuditService


class ExcelImportService:

    def __init__(self):
        self.compliance_service = ComplianceService()

    def process_excel(
        self,
        file_bytes: bytes,
        db: Session,
        dry_run: bool = True,
        user_id: str = "System"
    ) -> ImportExcelResponse:
        """Executes the complete safe Excel ingestion pipeline."""
        validation_errors: List[str] = []
        rows_parsed: Dict[str, int] = {
            "requirements": 0,
            "risks": 0,
            "controls": 0,
            "verifications": 0,
            "evidence": 0
        }
        persisted_counts: Dict[str, int] = {
            "requirements": 0,
            "risks": 0,
            "controls": 0,
            "verifications": 0,
            "evidence": 0
        }

        # Step 1: Open Workbook
        try:
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        except Exception as e:
            return ImportExcelResponse(
                success=False,
                dry_run=dry_run,
                message=f"Failed to open Excel workbook: {str(e)}",
                rows_parsed=rows_parsed,
                validation_errors=[f"Invalid or corrupted Excel file: {str(e)}"],
                compliance_findings=[],
                persisted_counts=persisted_counts
            )

        # Step 2: Validate Required Sheets
        required_sheets = ["Requirements", "Risks", "Controls", "Verifications", "Evidence"]
        for s in required_sheets:
            if s not in wb.sheetnames:
                validation_errors.append(f"Missing mandatory sheet: '{s}'")

        if validation_errors:
            return ImportExcelResponse(
                success=False,
                dry_run=dry_run,
                message="Workbook structure validation failed.",
                rows_parsed=rows_parsed,
                validation_errors=validation_errors,
                compliance_findings=[],
                persisted_counts=persisted_counts
            )

        # Step 3: Load into In-Memory DigitalThread
        thread = DigitalThread()
        try:
            # 1. Requirements
            ws_req = wb["Requirements"]
            headers_req = [c.value for c in ws_req[1] if c.value is not None]
            for row in ws_req.iter_rows(min_row=2, values_only=True):
                if not row or not row[0]:
                    continue
                d = dict(zip(headers_req, row))
                req_id = str(d["Req_ID"]).strip()
                from datetime import timezone
                created_dt = self.compliance_service._parse_iso_date(d.get("Created_Date")) or datetime.now(timezone.utc)
                mod_dt = self.compliance_service._parse_iso_date(d.get("Last_Modified_Date")) or datetime.now(timezone.utc)
                
                # Check description length
                desc = str(d.get("Description", "") or "")
                if len(desc) < 20:
                    validation_errors.append(f"Requirement '{req_id}': Description must be at least 20 characters.")

                thread.requirements[req_id] = self._make_req(d, req_id, created_dt, mod_dt)
                rows_parsed["requirements"] += 1

            # 2. Risks
            ws_rsk = wb["Risks"]
            headers_rsk = [c.value for c in ws_rsk[1] if c.value is not None]
            for row in ws_rsk.iter_rows(min_row=2, values_only=True):
                if not row or not row[0]:
                    continue
                d = dict(zip(headers_rsk, row))
                risk_id = str(d["Risk_ID"]).strip()
                parent_req_id = str(d.get("Parent_Req_ID", "")).strip()

                if parent_req_id not in thread.requirements:
                    # check if it already exists in database if not in thread
                    from backend.app.repositories.requirement_repo import RequirementRepository
                    if not RequirementRepository.get_by_id(db, parent_req_id):
                        validation_errors.append(f"Risk '{risk_id}': Parent requirement '{parent_req_id}' not found.")

                thread.risks[risk_id] = self._make_risk(d, risk_id, parent_req_id)
                rows_parsed["risks"] += 1

            # 3. Controls
            ws_ctrl = wb["Controls"]
            headers_ctrl = [c.value for c in ws_ctrl[1] if c.value is not None]
            for row in ws_ctrl.iter_rows(min_row=2, values_only=True):
                if not row or not row[0]:
                    continue
                d = dict(zip(headers_ctrl, row))
                ctrl_id = str(d["Control_ID"]).strip()
                parent_risk_id = str(d.get("Parent_Risk_ID", "")).strip()

                if parent_risk_id not in thread.risks:
                    from backend.app.repositories.risk_repo import RiskRepository
                    if not RiskRepository.get_by_id(db, parent_risk_id):
                        validation_errors.append(f"Control '{ctrl_id}': Parent risk '{parent_risk_id}' not found.")

                thread.controls[ctrl_id] = self._make_control(d, ctrl_id, parent_risk_id)
                rows_parsed["controls"] += 1

            # 4. Verifications
            ws_vrf = wb["Verifications"]
            headers_vrf = [c.value for c in ws_vrf[1] if c.value is not None]
            for row in ws_vrf.iter_rows(min_row=2, values_only=True):
                if not row or not row[0]:
                    continue
                d = dict(zip(headers_vrf, row))
                verif_id = str(d["Verif_ID"]).strip()
                target_req = str(d.get("Target_Req_ID")).strip() if d.get("Target_Req_ID") else None
                target_ctrl = str(d.get("Target_Control_ID")).strip() if d.get("Target_Control_ID") else None

                # Dual-target validation
                if (target_req and target_ctrl) or (not target_req and not target_ctrl):
                    validation_errors.append(
                        f"Verification '{verif_id}': Dual-target violation. Exactly one target must be populated."
                    )

                thread.verifications[verif_id] = self._make_verification(d, verif_id, target_req, target_ctrl)
                rows_parsed["verifications"] += 1

            # 5. Evidence
            ws_evd = wb["Evidence"]
            headers_evd = [c.value for c in ws_evd[1] if c.value is not None]
            for row in ws_evd.iter_rows(min_row=2, values_only=True):
                if not row or not row[0]:
                    continue
                d = dict(zip(headers_evd, row))
                evid_id = str(d["Evid_ID"]).strip()
                parent_verif = str(d.get("Parent_Verif_ID", "")).strip()

                if parent_verif not in thread.verifications:
                    from backend.app.repositories.verification_repo import VerificationRepository
                    if not VerificationRepository.get_by_id(db, parent_verif):
                        validation_errors.append(f"Evidence '{evid_id}': Parent verification '{parent_verif}' not found.")

                thread.evidence[evid_id] = self._make_evidence(d, evid_id, parent_verif)
                rows_parsed["evidence"] += 1

        except Exception as e:
            validation_errors.append(f"Normalization error: {str(e)}")

        if validation_errors:
            return ImportExcelResponse(
                success=False,
                dry_run=dry_run,
                message=f"Workbook parsing failed with {len(validation_errors)} error(s).",
                rows_parsed=rows_parsed,
                validation_errors=validation_errors,
                compliance_findings=[],
                persisted_counts=persisted_counts
            )

        # Step 4: Run Deterministic Compliance Engine
        eval_result = self.compliance_service.evaluate_thread(thread)
        findings_response = [
            FindingResponse(
                finding_id=f.finding_id,
                target_entity_type=f.target_entity_type,
                target_entity_id=f.target_entity_id,
                rule_violated=f.rule_violated,
                severity=f.severity,
                message=f.message,
                finding_status=f.finding_status,
                created_date=f.created_date or datetime.utcnow(),
                resolved_date=f.resolved_date,
                resolution=f.resolution,
                assigned_to=f.assigned_to
            )
            for f in eval_result.findings
        ]

        # Step 5: Dry-Run Check
        if dry_run:
            return ImportExcelResponse(
                success=True,
                dry_run=True,
                message="Dry-run completed successfully. Digital thread validated against deterministic engine.",
                rows_parsed=rows_parsed,
                validation_errors=[],
                compliance_findings=findings_response,
                persisted_counts=persisted_counts
            )

        # Step 6: Explicit Persistence Step (when dry_run=False)
        try:
            # Persist Requirements
            for r in thread.requirements.values():
                db_req = DBRequirement(
                    req_id=r.req_id,
                    version=r.version,
                    description=r.description,
                    category=r.category,
                    rationale=r.rationale,
                    source=r.source,
                    priority=r.priority,
                    owner=r.owner,
                    verification_method=r.verification_method,
                    acceptance_criteria=r.acceptance_criteria,
                    status=r.status,
                    created_date=r.created_date or datetime.utcnow(),
                    last_modified_date=r.last_modified_date or datetime.utcnow()
                )
                db.merge(db_req)
                persisted_counts["requirements"] += 1
                AuditService.log_mutation(
                    db=db,
                    entity_type="Requirement",
                    entity_id=r.req_id,
                    action=AuditAction.Create.value,
                    new_value={"req_id": r.req_id, "imported_from_excel": True},
                    user_id=user_id,
                    reason="Excel batch import"
                )

            # Persist Risks
            for rk in thread.risks.values():
                db_risk = DBRisk(
                    risk_id=rk.risk_id,
                    parent_req_id=rk.parent_req_id,
                    failure_mode=rk.failure_mode,
                    cause=rk.cause,
                    effect=rk.effect,
                    pre_severity=rk.pre_severity,
                    pre_likelihood=rk.pre_likelihood,
                    detectability=rk.detectability,
                    risk_category=rk.risk_category,
                    owner=rk.owner,
                    status=rk.status
                )
                db.merge(db_risk)
                persisted_counts["risks"] += 1
                AuditService.log_mutation(
                    db=db,
                    entity_type="Risk",
                    entity_id=rk.risk_id,
                    action=AuditAction.Create.value,
                    new_value={"risk_id": rk.risk_id, "imported_from_excel": True},
                    user_id=user_id,
                    reason="Excel batch import"
                )

            # Persist Controls
            for c in thread.controls.values():
                db_ctrl = DBControl(
                    control_id=c.control_id,
                    parent_risk_id=c.parent_risk_id,
                    hierarchy_type=c.hierarchy_type,
                    description=c.description,
                    rationale=c.rationale,
                    post_severity=c.post_severity,
                    post_likelihood=c.post_likelihood,
                    owner=c.owner,
                    status=c.status,
                    implementation_date=c.implementation_date
                )
                db.merge(db_ctrl)
                persisted_counts["controls"] += 1
                AuditService.log_mutation(
                    db=db,
                    entity_type="Control",
                    entity_id=c.control_id,
                    action=AuditAction.Create.value,
                    new_value={"control_id": c.control_id, "imported_from_excel": True},
                    user_id=user_id,
                    reason="Excel batch import"
                )

            # Persist Verifications
            for v in thread.verifications.values():
                db_v = DBVerification(
                    verif_id=v.verif_id,
                    target_type=v.target_type,
                    target_req_id=v.target_req_id,
                    target_control_id=v.target_control_id,
                    method=v.method,
                    acceptance_crit=v.acceptance_crit,
                    reviewer=v.reviewer,
                    status=v.status
                )
                db.merge(db_v)
                persisted_counts["verifications"] += 1
                AuditService.log_mutation(
                    db=db,
                    entity_type="Verification",
                    entity_id=v.verif_id,
                    action=AuditAction.Create.value,
                    new_value={"verif_id": v.verif_id, "imported_from_excel": True},
                    user_id=user_id,
                    reason="Excel batch import"
                )

            # Persist Evidence
            for e in thread.evidence.values():
                db_e = DBEvidence(
                    evid_id=e.evid_id,
                    parent_verif_id=e.parent_verif_id,
                    evidence_type=e.evidence_type,
                    file_name=e.file_name,
                    artifact_ref=e.artifact_ref,
                    hash=e.hash,
                    result_value=e.result_value,
                    date_generated=e.date_generated or datetime.utcnow(),
                    approval_status=e.approval_status
                )
                db.merge(db_e)
                persisted_counts["evidence"] += 1
                AuditService.log_mutation(
                    db=db,
                    entity_type="Evidence",
                    entity_id=e.evid_id,
                    action=AuditAction.Create.value,
                    new_value={"evid_id": e.evid_id, "imported_from_excel": True},
                    user_id=user_id,
                    reason="Excel batch import"
                )

            db.commit()
            return ImportExcelResponse(
                success=True,
                dry_run=False,
                message="Successfully validated and persisted workbook data into PostgreSQL.",
                rows_parsed=rows_parsed,
                validation_errors=[],
                compliance_findings=findings_response,
                persisted_counts=persisted_counts
            )
        except IntegrityError as ex:
            db.rollback()
            return ImportExcelResponse(
                success=False,
                dry_run=False,
                message=f"Database persistence error: {str(ex.orig) if hasattr(ex, 'orig') else str(ex)}",
                rows_parsed=rows_parsed,
                validation_errors=[f"Database constraint violation: {str(ex)}"],
                compliance_findings=findings_response,
                persisted_counts={"requirements": 0, "risks": 0, "controls": 0, "verifications": 0, "evidence": 0}
            )

    @staticmethod
    def _make_req(d: Dict[str, Any], req_id: str, c_dt: datetime, m_dt: datetime):
        from backend.app.domain.models import Requirement
        return Requirement(
            req_id=req_id,
            version=str(d.get("Version", "v1.0")),
            description=str(d.get("Description", "")),
            category=str(d.get("Category", "Structural_Loading")),
            rationale=str(d.get("Rationale", "") or ""),
            source=str(d.get("Source", "")),
            priority=str(d.get("Priority", "Must_Have")),
            owner=str(d.get("Owner")) if d.get("Owner") else None,
            verification_method=str(d.get("Verification_Method", "Analysis")),
            acceptance_criteria=str(d.get("Acceptance_Criteria", "")),
            status=str(d.get("Status", "Draft")),
            created_date=c_dt,
            last_modified_date=m_dt
        )

    @staticmethod
    def _make_risk(d: Dict[str, Any], risk_id: str, parent_req_id: str):
        from backend.app.domain.models import Risk
        pre_sev = int(d.get("Pre_Severity", 1) or 1)
        pre_like = int(d.get("Pre_Likelihood", 1) or 1)
        det = int(d.get("Detectability")) if d.get("Detectability") else None
        return Risk(
            risk_id=risk_id,
            parent_req_id=parent_req_id,
            failure_mode=str(d.get("Failure_Mode", "")),
            cause=str(d.get("Cause", "")),
            effect=str(d.get("Effect", "")),
            pre_severity=pre_sev,
            pre_likelihood=pre_like,
            detectability=det,
            risk_score=pre_sev * pre_like,
            risk_category=str(d.get("Risk_Category", "Structural_Yielding")),
            owner=str(d.get("Owner")) if d.get("Owner") else None,
            status=str(d.get("Status", "Identified"))
        )

    @staticmethod
    def _make_control(d: Dict[str, Any], ctrl_id: str, parent_risk_id: str):
        from backend.app.domain.models import Control
        post_sev = int(d.get("Post_Severity", 1) or 1)
        post_like = int(d.get("Post_Likelihood", 1) or 1)
        return Control(
            control_id=ctrl_id,
            parent_risk_id=parent_risk_id,
            hierarchy_type=str(d.get("Hierarchy_Type", "Inherently_Safe_Design")),
            description=str(d.get("Description", "")),
            rationale=str(d.get("Rationale")) if d.get("Rationale") else None,
            post_severity=post_sev,
            post_likelihood=post_like,
            residual_risk=post_sev * post_like,
            owner=str(d.get("Owner")) if d.get("Owner") else None,
            status=str(d.get("Status", "Proposed")),
            implementation_date=None
        )

    @staticmethod
    def _make_verification(d: Dict[str, Any], verif_id: str, target_req: Optional[str], target_ctrl: Optional[str]):
        from backend.app.domain.models import Verification
        target_type = "Requirement" if target_req else "Control"
        return Verification(
            verif_id=verif_id,
            target_type=target_type,
            target_req_id=target_req,
            target_control_id=target_ctrl,
            method=str(d.get("Method", "Analysis")),
            acceptance_crit=str(d.get("Acceptance_Crit", "")),
            reviewer=str(d.get("Reviewer")) if d.get("Reviewer") else None,
            status=str(d.get("Status", "Planned"))
        )

    def _make_evidence(self, d: Dict[str, Any], evid_id: str, parent_verif: str):
        from backend.app.domain.models import Evidence
        from datetime import timezone
        gen_dt = self.compliance_service._parse_iso_date(d.get("Date_Generated")) or datetime.now(timezone.utc)
        return Evidence(
            evid_id=evid_id,
            parent_verif_id=parent_verif,
            evidence_type=str(d.get("Evidence_Type", "Simulation_Report")),
            file_name=str(d.get("File_Name", "")),
            artifact_ref=str(d.get("Artifact_Ref", "")),
            hash=str(d.get("Hash")) if d.get("Hash") else None,
            result_value=str(d.get("Result_Value")) if d.get("Result_Value") else None,
            date_generated=gen_dt,
            approval_status=str(d.get("Approval_Status", "Uploaded"))
        )
